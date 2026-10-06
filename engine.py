import os
import subprocess
from PIL import Image
import pdf2docx
import pdfplumber
import pandas as pd
from pptx import Presentation
from pptx.util import Inches

class DocumentEngine:
    @staticmethod
    def docx_to_pdf(input_path: str, output_dir: str) -> str:
        """Converts DOCX, XLSX, or PPTX to PDF using headless LibreOffice."""
        cmd = ["soffice", "--headless", "--convert-to", "pdf", "--outdir", output_dir, input_path]
        subprocess.run(cmd, check=True)
        base_name = os.path.splitext(os.path.basename(input_path))[0]
        filename = f"{base_name}.pdf"
        return os.path.join(output_dir, filename)

    @staticmethod
    def pdf_to_docx(input_path: str, output_path: str):
        """Converts PDF to editable Word DOCX preserving layout."""
        cv = pdf2docx.Converter(input_path)
        cv.convert(output_path, start=0, end=None)
        cv.close()
        return output_path

    @staticmethod
    def pdf_to_excel(input_path: str, output_path: str):
        """Extracts tables from PDF and writes them to Excel (.xlsx)."""
        all_tables = []
        with pdfplumber.open(input_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    if table and len(table) > 1:
                        headers = [str(h) if h is not None else f"Col_{i}" for i, h in enumerate(table[0])]
                        df = pd.DataFrame(table[1:], columns=headers)
                        all_tables.append(df)
        
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            if all_tables:
                for idx, df in enumerate(all_tables):
                    df.to_excel(writer, sheet_name=f"Table_{idx+1}", index=False)
            else:
                pd.DataFrame({"Note": ["No structural tables detected"]}).to_excel(writer, sheet_name="Sheet1", index=False)
        return output_path

    @staticmethod
    def pdf_to_pptx(input_path: str, output_path: str):
        """Converts PDF pages into high-resolution PowerPoint slides."""
        from pdf2image import convert_from_path
        images = convert_from_path(input_path, dpi=200)
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]

        for i, img in enumerate(images):
            temp_img_path = f"temp_page_{i}.png"
            img.save(temp_img_path, "PNG")
            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(temp_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
            if os.path.exists(temp_img_path):
                os.remove(temp_img_path)

        prs.save(output_path)
        return output_path

class CompressionEngine:
    @staticmethod
    def compress_pdf(input_path: str, output_path: str, preset: str = "ebook") -> str:
        """Compresses PDF using Ghostscript."""
        gs_presets = {
            "low": "/screen",
            "medium": "/ebook",
            "high": "/printer"
        }
        pdf_setting = gs_presets.get(preset, "/ebook")
        gs_cmd = "gswin64c" if os.name == "nt" else "gs"
        cmd = [
            gs_cmd, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4",
            f"-dPDFSETTINGS={pdf_setting}", "-dNOPAUSE", "-dQUIET", "-dBATCH",
            f"-sOutputFile={output_path}", input_path
        ]
        subprocess.run(cmd, check=True)
        return output_path

    @staticmethod
    def resize_image(input_path: str, output_path: str, max_width: int = 1920, quality: int = 80) -> str:
        """Resizes image dimensions and compresses quality."""
        with Image.open(input_path) as img:
            img = img.convert("RGB")
            if img.width > max_width:
                ratio = max_width / float(img.width)
                new_height = int(float(img.height) * float(ratio))
                img = img.resize((max_width, new_height), Image.Resampling.LANCZOS)
            
            img.save(output_path, "JPEG", optimize=True, quality=quality)
        return output_path
